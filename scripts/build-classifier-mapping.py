#!/usr/bin/env python3
"""Build normalized classifier applicability mapping from the approved XLSX."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import openpyxl


SPACE = re.compile(r"\s+")
EXPECTED_CODES = {
    *(f"1.{index}" for index in range(1, 9)),
    *(f"2.{index}" for index in range(1, 7)),
    *(f"3.{index}" for index in range(1, 10)),
    *(f"4.{index}" for index in range(1, 14)),
    *(f"5.{index}" for index in range(1, 5)),
    *(f"6.{index}" for index in range(1, 3)),
    *(f"7.{index}" for index in range(1, 8)),
}


def clean(value: object) -> str:
    return "" if value is None else SPACE.sub(" ", str(value)).strip()


def values(value: object) -> list[str]:
    raw = clean(value)
    if not raw or raw in {"—", "-"}:
        return []
    return list(dict.fromkeys(part.strip() for part in raw.split(",") if part.strip()))


def criterion_code(value: object, name: str) -> str:
    code = f"{float(value):.2f}".rstrip("0").rstrip(".")
    if code == "4.1" and "лиценз" in name.casefold():
        return "4.10"
    return code


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    sheet = openpyxl.load_workbook(args.source, read_only=True, data_only=True)["Mapping"]
    rows: list[dict[str, object]] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        name = clean(row[2])
        code = criterion_code(row[1], name)
        rows.append({
            "code": code,
            "section": clean(row[0]),
            "name": name,
            "categories": values(row[3]),
            "types": values(row[4]),
            "zones": values(row[5]),
            "equipment": values(row[6]),
            "specialZones": values(row[7]),
            "environmentalAspects": values(row[8]),
            "regions": values(row[9]),
        })

    actual_codes = {str(row["code"]) for row in rows}
    if len(rows) != 49 or actual_codes != EXPECTED_CODES:
        raise ValueError(f"Mapping must contain 49 unique classifier codes; missing={sorted(EXPECTED_CODES-actual_codes)}, extra={sorted(actual_codes-EXPECTED_CODES)}")

    rows.sort(key=lambda row: tuple(map(int, str(row["code"]).split("."))))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="\n") as output:
        for row in rows:
            output.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")

    manifest = {
        "sourceFile": args.source.name,
        "sourceSha256": hashlib.sha256(args.source.read_bytes()).hexdigest(),
        "sourceRows": sheet.max_row - 1,
        "mappingEntries": len(rows),
        "correctedSourceCodes": {"licensingCriterion": "4.1 -> 4.10"},
    }
    args.output.with_name(f"{args.output.stem}.manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Classifier mapping built: {len(rows)} criteria.")


if __name__ == "__main__":
    main()
