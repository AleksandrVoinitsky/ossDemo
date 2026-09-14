#!/usr/bin/env python3
"""Build the deterministic requirements catalog from the approved XLSX registry."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

import openpyxl


CODE_PATTERN = re.compile(r"(?<!\d)([1-7]\.\d{1,2})(?!\d)")
WHITESPACE_PATTERN = re.compile(r"\s+")
ORPHAN_CODES = {"2.7"}


def text(value: object) -> str:
    if value is None:
        return ""
    return WHITESPACE_PATTERN.sub(" ", str(value)).strip()


def version(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return text(value)


def append_unique(target: list[str], value: str) -> None:
    if value and value.casefold() not in {item.casefold() for item in target}:
        target.append(value)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.source, read_only=True, data_only=True)
    sheet = workbook["Реестр требований"]
    catalog: OrderedDict[tuple[str, str], dict[str, object]] = OrderedDict()

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        basis = text(row[4])
        requirement = text(row[5])
        if not basis or not requirement:
            raise ValueError(f"Row {row_number}: basis and requirement are required")
        codes = [code for code in CODE_PATTERN.findall(text(row[11])) if code not in ORPHAN_CODES]
        if not codes:
            raise ValueError(f"Row {row_number}: no supported classifier code")

        key = (basis.casefold(), requirement.casefold())
        entry = catalog.get(key)
        if entry is None:
            digest = hashlib.sha256(f"{basis}\0{requirement}".encode("utf-8")).hexdigest()[:24]
            entry = {
                "id": digest,
                "version": version(row[0]),
                "levels": [],
                "groups": [],
                "classifierCodes": [],
                "basis": basis,
                "requirement": requirement,
                "categories": [],
                "liability": [],
                "penalties": [],
                "notes": [],
            }
            catalog[key] = entry

        for field, value in (
            ("levels", text(row[1])),
            ("groups", text(row[10]) or text(row[2])),
            ("categories", text(row[8])),
            ("liability", text(row[6])),
            ("penalties", text(row[7])),
            ("notes", text(row[9])),
        ):
            append_unique(entry[field], value)  # type: ignore[arg-type]
        for code in codes:
            append_unique(entry["classifierCodes"], code)  # type: ignore[arg-type]

    entries = sorted(
        catalog.values(),
        key=lambda item: (
            min((tuple(map(int, code.split("."))) for code in item["classifierCodes"]), default=(99, 99)),
            str(item["basis"]).casefold(),
            str(item["requirement"]).casefold(),
        ),
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="\n") as output:
        for entry in entries:
            output.write(json.dumps(entry, ensure_ascii=False, separators=(",", ":")) + "\n")

    manifest = {
        "sourceFile": args.source.name,
        "sourceSha256": hashlib.sha256(args.source.read_bytes()).hexdigest(),
        "sourceRows": sheet.max_row - 1,
        "catalogEntries": len(entries),
        "excludedClassifierCodes": sorted(ORPHAN_CODES),
    }
    manifest_path = args.output.with_name(f"{args.output.stem}.manifest.json")
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"Catalog built: {len(entries)} unique requirements from {sheet.max_row - 1} rows.")


if __name__ == "__main__":
    main()
