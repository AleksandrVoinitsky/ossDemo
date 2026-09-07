"""Convert archived environmental inspection XLSX checklists to UTF-8 CSV data files.

The CSV files are source data for the application checklist service. They are
intentionally independent from the RAG knowledge-base import pipeline.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIRECTORY = ROOT / "02_📂 Чек-листы" / "Исторические чек-листы"
DESTINATION_DIRECTORY = ROOT / "src" / "OssDemo.Web" / "Data" / "Checklists"


def text(value: object | None) -> str:
    return "" if value is None else " ".join(str(value).split())


def slug(value: str) -> str:
    transliteration = str.maketrans(
        "абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
        "abvgdeejzijklmnoprstufhccss_y_eua",
    )
    return re.sub(r"[^a-z0-9]+", "-", value.lower().translate(transliteration)).strip("-")


def parse_file(path: Path) -> tuple[dict[str, str], list[dict[str, str]]]:
    worksheet = load_workbook(path, read_only=True, data_only=True).active
    rows = list(worksheet.iter_rows(values_only=True))
    facility = next((text(row[0]).split(":", 1)[-1].strip() for row in rows if text(row[0]).startswith("Проверяемое структурное подразделение")), "")
    period = next((text(row[0]).split(":", 1)[-1].strip() for row in rows if text(row[0]).startswith("Период проведения")), "")
    inspector = next((text(row[0]).split(":", 1)[-1].strip() for row in rows if text(row[0]).startswith("ФИО, должность")), "")
    items: list[dict[str, str]] = []
    section = "Общие вопросы"

    for row in rows:
        number = row[0] if row else None
        if isinstance(number, (int, float)) and len(row) >= 6 and text(row[1]):
            items.append(
                {
                    "number": str(int(number)),
                    "section": section,
                    "title": text(row[1]),
                    "basis": text(row[2]),
                    "result": text(row[3]),
                    "nonconformity": text(row[4]),
                    "note": text(row[5]),
                }
            )
        elif text(number) and not text(row[1]) and not text(row[2]) and "Чек-лист" not in text(number) and "Проверяемое" not in text(number):
            section = text(number)

    return {"facility": facility, "period": period, "inspector": inspector}, items


def main() -> None:
    DESTINATION_DIRECTORY.mkdir(parents=True, exist_ok=True)
    for source_path in sorted(SOURCE_DIRECTORY.glob("*.xlsx")):
        metadata, items = parse_file(source_path)
        output_path = DESTINATION_DIRECTORY / f"{slug(source_path.stem)}.csv"
        with output_path.open("w", encoding="utf-8-sig", newline="") as output:
            writer = csv.DictWriter(
                output,
                fieldnames=["number", "section", "title", "basis", "result", "nonconformity", "note"],
            )
            writer.writeheader()
            writer.writerows(items)
        print(f"{source_path.name}: {len(items)} пунктов -> {output_path.name}")
        print(f"  Объект: {metadata['facility']}; период: {metadata['period']}; инспектор: {metadata['inspector']}")


if __name__ == "__main__":
    main()
